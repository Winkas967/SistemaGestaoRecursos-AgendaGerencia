from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.evaluations_service import EvaluationService


#gera a planilha excel com os indicadores do dashboard de avaliacoes,
#para consulta/edicao facil fora do sistema
class DashboardExportService:

    HEADER_FILL = PatternFill(start_color="FF1E7A4C", end_color="FF1E7A4C", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFFFF", bold=True)
    TOTAL_FILL = PatternFill(start_color="FFE8F3EC", end_color="FFE8F3EC", fill_type="solid")
    TOTAL_FONT = Font(bold=True)

    #escreve uma linha de cabecalho estilizada na planilha
    @staticmethod
    def _write_header(sheet, headers):
        sheet.append(headers)
        for column_index in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=column_index)
            cell.fill = DashboardExportService.HEADER_FILL
            cell.font = DashboardExportService.HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


    #ajusta a largura das colunas com base no conteudo
    @staticmethod
    def _autosize_columns(sheet):
        for column_cells in sheet.columns:
            length = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=8,
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 42)


    #monta a aba com o resumo agregado por categoria
    @staticmethod
    def _build_summary_sheet(workbook, dashboard):
        sheet = workbook.active
        sheet.title = "Resumo por categoria"

        headers = [
            "Categoria",
            "Prestadores",
            "Adesão",
            "Não adesão",
            "Não se posicionaram",
            "Sem visita",
            "Atendimento Centro médico/EVB",
            "Visita sem documento",
            "05 estrelas",
            "04 estrelas",
            "03 estrelas",
            "02 estrelas",
            "01 estrela",
            "00 estrelas",
        ]
        DashboardExportService._write_header(sheet, headers)

        for item in dashboard["categorias"]:
            sheet.append([
                item["categoriaNome"],
                item["totalPrestadores"],
                item["adesao"],
                item["naoAdesao"],
                item["naoPosicionaram"],
                item["semVisita"],
                item["atendimentoCentroMedico"],
                item["visitaSemDocumento"],
                item["estrelas5"],
                item["estrelas4"],
                item["estrelas3"],
                item["estrelas2"],
                item["estrelas1"],
                item["estrelas0"],
            ])

        totais = dashboard["totais"]
        total_row = [
            "TOTAL",
            totais["totalPrestadores"],
            totais["adesao"],
            totais["naoAdesao"],
            totais["naoPosicionaram"],
            totais["semVisita"],
            totais["atendimentoCentroMedico"],
            totais["visitaSemDocumento"],
            totais["estrelas5"],
            totais["estrelas4"],
            totais["estrelas3"],
            totais["estrelas2"],
            totais["estrelas1"],
            totais["estrelas0"],
        ]
        sheet.append(total_row)
        last_row = sheet.max_row
        for column_index in range(1, len(headers) + 1):
            cell = sheet.cell(row=last_row, column=column_index)
            cell.fill = DashboardExportService.TOTAL_FILL
            cell.font = DashboardExportService.TOTAL_FONT

        sheet.append([])
        sheet.append(["Média de adesão", f"{dashboard['mediaAdesaoPercentual']}%"])

        DashboardExportService._autosize_columns(sheet)


    #monta a aba com o detalhamento linha a linha por prestador
    @staticmethod
    def _build_details_sheet(workbook, details):
        sheet = workbook.create_sheet("Detalhado por prestador")

        headers = [
            "Categoria",
            "Prestador",
            "Status de adesão",
            "Status da avaliação",
            "Teve visita",
            "Estrelas",
            "Resultado (%)",
            "Iniciado em",
            "Concluído em",
        ]
        DashboardExportService._write_header(sheet, headers)

        for item in details["registros"]:
            teve_visita = item["teveVisita"]
            teve_visita_label = "—" if teve_visita is None else ("Sim" if teve_visita else "Não")

            sheet.append([
                item["categoriaNome"],
                item["prestadorNome"],
                item["statusAdesao"],
                item["statusAvaliacao"],
                teve_visita_label,
                item["estrelas"] if item["estrelas"] is not None else "—",
                item["resultadoPercentual"] if item["resultadoPercentual"] is not None else "—",
                item["iniciadoEm"].strftime("%d/%m/%Y %H:%M") if item["iniciadoEm"] else "—",
                item["concluidoEm"].strftime("%d/%m/%Y %H:%M") if item["concluidoEm"] else "—",
            ])

        DashboardExportService._autosize_columns(sheet)


    #gera o arquivo .xlsx completo do dashboard para o ano informado
    @staticmethod
    def build_workbook(year=None):
        dashboard = EvaluationService.get_dashboard(year)
        details = EvaluationService.get_dashboard_details_completo(dashboard["anoReferencia"])

        workbook = Workbook()
        DashboardExportService._build_summary_sheet(workbook, dashboard)
        DashboardExportService._build_details_sheet(workbook, details)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        return buffer, dashboard["anoReferencia"]
